"""گزارش‌های مالی: درآمد، سود واقعی، Ledger، جستجوی تراکنش، مرکز بازپرداخت، پرداخت‌های ناموفق،
ماشین‌حساب کارمزد و خروجی CSV / Excel / PDF.

تعریف‌ها (همه بدون سفارش‌های Test Mode):
  درآمد ناخالص  = مجموع مبلغ سفارش‌های انجام‌شده
  هزینه          = مجموع قیمت خرید از Stard همان سفارش‌ها
  کارمزدها       = درصدها/مبلغ ثابت تنظیم‌شده (تلگرام، API، درگاه پرداخت) روی درآمد
  پرداخت‌ها      = پاداش زیرمجموعه + پاداش روزانه/گردونه
  سود خالص       = درآمد − هزینه − کارمزدها − پرداخت‌ها
  برگشتی‌ها جدا نمایش داده می‌شوند (سفارش برگشتی درآمد حساب نمی‌شود چون پولش به کاربر برگشته).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import Any

from .db import Database, ts

DEFAULT_FEES = {"telegram_percent": 0.0, "api_percent": 0.0, "payment_percent": 0.0, "payment_fixed": 0}


def day_start(d: date) -> str:
    return f"{d.isoformat()}T00:00:00Z"


def utc_today() -> date:
    return datetime.now(timezone.utc).date()


def period_range(period: str, *, start: date | None = None, end: date | None = None) -> tuple[str, str, str]:
    """(از، تا، برچسب) برای day | week | month | custom. «تا» شامل نیست."""
    today = utc_today()
    if period == "day":
        a, b = today, today + timedelta(days=1)
    elif period == "week":
        a, b = today - timedelta(days=6), today + timedelta(days=1)
    elif period == "month":
        a, b = today - timedelta(days=29), today + timedelta(days=1)
    elif period == "custom" and start and end:
        a, b = start, end + timedelta(days=1)
    else:
        raise ValueError("invalid period")
    return day_start(a), day_start(b), f"{a.isoformat()} … {(b - timedelta(days=1)).isoformat()}"


async def fees(db: Database) -> dict:
    return {**DEFAULT_FEES, **(await db.get_json("fees", {}) or {})}


def fee_amount(revenue: int, orders: int, f: dict) -> int:
    pct = float(f["telegram_percent"]) + float(f["api_percent"]) + float(f["payment_percent"])
    return int(round(revenue * pct / 100 + orders * float(f["payment_fixed"])))


@dataclass
class Summary:
    revenue: int
    cost: int
    orders: int
    aov: int
    refunds: int
    refunded_orders: int
    payouts: int
    fees: int
    gross_profit: int
    net_profit: int
    margin: float


async def summary(db: Database, since: str, until: str) -> Summary:
    r = await db.one(
        """SELECT COALESCE(SUM(price), 0) AS revenue, COALESCE(SUM(base_amount), 0) AS cost, COUNT(*) AS orders
           FROM orders WHERE status = 'completed' AND is_test = 0 AND created_at >= :a AND created_at < :b""",
        {"a": since, "b": until})
    ref = await db.one("SELECT COALESCE(SUM(amount), 0) AS amount, COUNT(*) AS n FROM refunds WHERE status = "
                       "'completed' AND created_at >= :a AND created_at < :b", {"a": since, "b": until})
    pay = await db.scalar("SELECT COALESCE(SUM(amount), 0) FROM ledger WHERE kind IN ('referral', 'reward') "
                          "AND created_at >= :a AND created_at < :b", {"a": since, "b": until})
    f = fee_amount(r["revenue"], r["orders"], await fees(db))
    gross = r["revenue"] - r["cost"]
    net = gross - f - pay
    return Summary(revenue=r["revenue"], cost=r["cost"], orders=r["orders"],
                   aov=r["revenue"] // r["orders"] if r["orders"] else 0, refunds=ref["amount"],
                   refunded_orders=ref["n"], payouts=pay, fees=f, gross_profit=gross, net_profit=net,
                   margin=round(net / r["revenue"] * 100, 1) if r["revenue"] else 0.0)


def growth(cur: int, prev: int) -> str:
    if prev == 0:
        return "—" if cur == 0 else "جدید"
    g = (cur - prev) / prev * 100
    return f"{'▲' if g >= 0 else '▼'} {abs(g):.1f}%"


async def revenue_dashboard(db: Database) -> dict[str, Any]:
    today = utc_today()
    out: dict[str, Any] = {}
    for key, days in (("today", 1), ("week", 7), ("month", 30)):
        a = day_start(today - timedelta(days=days - 1))
        b = day_start(today + timedelta(days=1))
        pa = day_start(today - timedelta(days=2 * days - 1))
        cur, prev = await summary(db, a, b), await summary(db, pa, a)
        out[key] = {"cur": cur, "prev": prev, "growth": growth(cur.revenue, prev.revenue)}
    out["total"] = await summary(db, "0000", "9999")
    return out


async def series(db: Database, granularity: str, points: int = 14) -> list[tuple[str, int, int]]:
    """[(برچسب، درآمد، سود)] برای daily | weekly | monthly — قدیمی به جدید."""
    today = utc_today()
    buckets: list[tuple[str, date, date]] = []
    if granularity == "daily":
        for i in range(points - 1, -1, -1):
            d = today - timedelta(days=i)
            buckets.append((d.strftime("%m-%d"), d, d + timedelta(days=1)))
    elif granularity == "weekly":
        start = today - timedelta(days=today.weekday())
        for i in range(points - 1, -1, -1):
            a = start - timedelta(weeks=i)
            buckets.append((a.strftime("%m-%d"), a, a + timedelta(weeks=1)))
    elif granularity == "monthly":
        y, m = today.year, today.month
        months = []
        for _ in range(points):
            months.append((y, m))
            m -= 1
            if m == 0:
                y, m = y - 1, 12
        for y, m in reversed(months):
            a = date(y, m, 1)
            b = date(y + (m == 12), m % 12 + 1, 1)
            buckets.append((a.strftime("%Y-%m"), a, b))
    else:
        raise ValueError(granularity)
    out = []
    for label, a, b in buckets:
        r = await db.one("SELECT COALESCE(SUM(price), 0) AS rev, COALESCE(SUM(price - base_amount), 0) AS profit "
                         "FROM orders WHERE status = 'completed' AND is_test = 0 AND created_at >= :a AND created_at < :b",
                         {"a": day_start(a), "b": day_start(b)})
        out.append((label, r["rev"], r["profit"]))
    return out


def text_chart(points: list[tuple[str, int]], width: int = 14) -> str:
    """نمودار میله‌ای متنی (یک سری، بدون رنگ): برچسب، میله‌ی متناسب، و مقدار عددی کنار هر میله."""
    if not points:
        return "—"
    top = max(v for _, v in points) or 1
    lines = []
    for label, v in points:
        n = int(round(max(v, 0) / top * width))
        bar = "█" * n if n else ("▏" if v > 0 else "·")
        lines.append(f"<code>{label} {bar.ljust(width)}</code> {_short_amount(v)}")
    return "\n".join(lines)


def _short_amount(v: int) -> str:
    if abs(v) >= 1_000_000_000:
        return f"{v / 1_000_000_000:.1f}B"
    if abs(v) >= 1_000_000:
        return f"{v / 1_000_000:.1f}M"
    if abs(v) >= 1_000:
        return f"{v / 1_000:.0f}K"
    return str(v)


# ---------- Ledger و جستجو ----------
QUERY_KEYS = {"user", "type", "min", "max", "from", "to", "order", "id", "status", "ref"}
_TOKEN = re.compile(r"(\w+):(\S+)")


def parse_query(q: str) -> dict[str, str]:
    """«user:123 type:refund min:1000 from:2026-10-01 to:2026-10-05 order:42 id:15 status:completed»"""
    out = {k.lower(): v for k, v in _TOKEN.findall(q or "") if k.lower() in QUERY_KEYS}
    rest = _TOKEN.sub("", q or "").strip()
    if rest.isdigit() and not out:
        out["any"] = rest  # عدد تنها: شناسه‌ی کاربر، سفارش یا تراکنش
    return out


def _date(v: str, end: bool = False) -> str:
    d = datetime.strptime(v, "%Y-%m-%d").date()
    return day_start(d + timedelta(days=1)) if end else day_start(d)


async def ledger_search(db: Database, q: str, *, limit: int = 20, offset: int = 0) -> list[dict]:
    f = parse_query(q)
    where, p = ["1=1"], {"l": limit, "o": offset}
    if "any" in f:
        where.append("(user_id = :n OR id = :n OR ref = :ns)")
        p.update(n=int(f["any"]), ns=f["any"])
    if "user" in f:
        where.append("user_id = :u")
        p["u"] = int(f["user"])
    if "type" in f:
        where.append("kind = :k")
        p["k"] = f["type"]
    if "id" in f:
        where.append("id = :id")
        p["id"] = int(f["id"])
    if "order" in f:
        where.append("ref = :oref AND kind IN ('order', 'refund', 'referral')")
        p["oref"] = f["order"]
    if "ref" in f:
        where.append("ref = :ref")
        p["ref"] = f["ref"]
    if "min" in f:
        where.append("ABS(amount) >= :mn")
        p["mn"] = int(f["min"])
    if "max" in f:
        where.append("ABS(amount) <= :mx")
        p["mx"] = int(f["max"])
    if "from" in f:
        where.append("created_at >= :fr")
        p["fr"] = _date(f["from"])
    if "to" in f:
        where.append("created_at < :to")
        p["to"] = _date(f["to"], end=True)
    return await db.all(f"SELECT * FROM ledger WHERE {' AND '.join(where)} ORDER BY id DESC LIMIT :l OFFSET :o", p)


async def order_search(db: Database, q: str, *, limit: int = 15) -> list[dict]:
    f = parse_query(q)
    where, p = ["1=1"], {"l": limit}
    if "any" in f:
        where.append("(id = :n OR user_id = :n)")
        p["n"] = int(f["any"])
    if "order" in f or "id" in f:
        where.append("id = :oid")
        p["oid"] = int(f.get("order") or f["id"])
    if "user" in f:
        where.append("user_id = :u")
        p["u"] = int(f["user"])
    if "status" in f:
        where.append("status = :s")
        p["s"] = f["status"]
    if "min" in f:
        where.append("price >= :mn")
        p["mn"] = int(f["min"])
    if "max" in f:
        where.append("price <= :mx")
        p["mx"] = int(f["max"])
    if "from" in f:
        where.append("created_at >= :fr")
        p["fr"] = _date(f["from"])
    if "to" in f:
        where.append("created_at < :to")
        p["to"] = _date(f["to"], end=True)
    return await db.all(f"SELECT * FROM orders WHERE {' AND '.join(where)} ORDER BY id DESC LIMIT :l", p)


async def refunds(db: Database, status: str | None = None, limit: int = 20) -> list[dict]:
    if status:
        return await db.all("SELECT * FROM refunds WHERE status = :s ORDER BY updated_at DESC LIMIT :l",
                            {"s": status, "l": limit})
    return await db.all("SELECT * FROM refunds ORDER BY updated_at DESC LIMIT :l", {"l": limit})


async def refund_counts(db: Database) -> dict[str, dict]:
    rows = await db.all("SELECT status, COUNT(*) AS n, COALESCE(SUM(amount), 0) AS amount FROM refunds GROUP BY status")
    return {r["status"]: r for r in rows}


async def failed_payments(db: Database, limit: int = 20) -> list[dict]:
    """سفارش‌هایی که به خاطر خطای API برگشت خوردند یا هنوز ارسالشان ناموفق است، به‌علاوه‌ی کارهای dead."""
    orders = await db.all(
        "SELECT o.*, (SELECT j.id FROM jobs j WHERE j.status = 'dead' AND j.kind LIKE 'order.%' "
        "AND j.payload LIKE '%\"oid\": ' || o.id || '}%' ORDER BY j.id DESC LIMIT 1) AS dead_job "
        "FROM orders o WHERE (o.refunded = 1 AND o.failure_reason IS NOT NULL) OR "
        "(o.status = 'new' AND o.failure_reason LIKE 'retrying:%') ORDER BY o.id DESC LIMIT :l", {"l": limit})
    return orders


# ---------- ماشین‌حساب کارمزد ----------
def fee_breakdown(sale: int, cost: int, f: dict, refund_rate: float = 0.0) -> dict[str, Any]:
    tg = round(sale * float(f["telegram_percent"]) / 100)
    api = round(sale * float(f["api_percent"]) / 100)
    pay = round(sale * float(f["payment_percent"]) / 100 + float(f["payment_fixed"]))
    refund = round(sale * refund_rate / 100)
    net = sale - cost - tg - api - pay - refund
    return {"sale": sale, "cost": cost, "telegram": tg, "api": api, "payment": pay, "refund": refund,
            "net": net, "margin": round(net / sale * 100, 2) if sale else 0.0}


# ---------- خروجی ----------
ORDER_COLS = ["id", "created_at", "user_id", "category", "title", "quantity", "recipient", "status", "price",
              "base_amount", "discount", "coupon", "refunded", "stard_ref", "failure_reason"]
LEDGER_COLS = ["id", "created_at", "user_id", "kind", "amount", "ref"]


async def report_rows(db: Database, since: str, until: str) -> tuple[list[dict], list[dict]]:
    orders = await db.all(f"SELECT {', '.join(ORDER_COLS)} FROM orders WHERE is_test = 0 AND created_at >= :a "
                          "AND created_at < :b ORDER BY id", {"a": since, "b": until})
    ledger = await db.all(f"SELECT {', '.join(LEDGER_COLS)} FROM ledger WHERE created_at >= :a AND created_at < :b "
                          "ORDER BY id", {"a": since, "b": until})
    return orders, ledger


def _safe_cell(v: Any) -> Any:
    # جلوگیری از CSV/Formula Injection در Excel (مقداری که با = + - @ شروع شود)
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


async def export_csv(db: Database, since: str, until: str) -> bytes:
    orders, ledger = await report_rows(db, since, until)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["# orders"])
    w.writerow(ORDER_COLS)
    for o in orders:
        w.writerow([_safe_cell(o[c]) for c in ORDER_COLS])
    w.writerow([])
    w.writerow(["# ledger"])
    w.writerow(LEDGER_COLS)
    for r in ledger:
        w.writerow([_safe_cell(r[c]) for c in LEDGER_COLS])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM برای نمایش درست فارسی در Excel


async def export_xlsx(db: Database, since: str, until: str, label: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    s = await summary(db, since, until)
    orders, ledger = await report_rows(db, since, until)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Stard Shop Bot — Financial report", label])
    ws["A1"].font = Font(bold=True, size=13)
    for k, v in (("Revenue", s.revenue), ("Cost", s.cost), ("Fees", s.fees), ("Payouts", s.payouts),
                 ("Gross profit", s.gross_profit), ("Net profit", s.net_profit), ("Margin %", s.margin),
                 ("Orders", s.orders), ("Average order value", s.aov), ("Refunds", s.refunds),
                 ("Refunded orders", s.refunded_orders)):
        ws.append([k, v])
    for title, cols, rows in (("Orders", ORDER_COLS, orders), ("Ledger", LEDGER_COLS, ledger)):
        sh = wb.create_sheet(title)
        sh.append(cols)
        for c in sh[1]:
            c.font = Font(bold=True)
        for r in rows:
            sh.append([_safe_cell(r[c]) for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_pdf(db: Database, since: str, until: str, label: str) -> bytes:
    """PDF خلاصه (برچسب‌ها انگلیسی‌اند چون فونت‌های داخلی PDF حروف فارسی ندارند)."""
    from fpdf import FPDF
    s = await summary(db, since, until)
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Stard Shop Bot - Financial report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    safe = label.replace("…", "to").encode("latin-1", "replace").decode("latin-1")
    pdf.cell(0, 8, f"Period: {safe}   Generated: {ts(datetime.now(timezone.utc))}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    rows = [("Revenue", s.revenue), ("Cost (Stard)", s.cost), ("Fees", s.fees), ("Payouts (referral/rewards)", s.payouts),
            ("Gross profit", s.gross_profit), ("Net profit", s.net_profit), ("Margin", f"{s.margin}%"),
            ("Completed orders", s.orders), ("Average order value", s.aov), ("Refunds (amount)", s.refunds),
            ("Refunded orders", s.refunded_orders)]
    for k, v in rows:
        pdf.set_font("Helvetica", "B" if k == "Net profit" else "", 11)
        pdf.cell(90, 8, k, border=1)
        pdf.cell(70, 8, f"{v:,} IRT" if isinstance(v, int) and k not in ("Completed orders", "Refunded orders")
                 else str(v), border=1, align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Daily revenue (last 14 days)", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for lbl, rev, profit in await series(db, "daily", 14):
        pdf.cell(30, 6, lbl, border=1)
        pdf.cell(50, 6, f"{rev:,}", border=1, align="R")
        pdf.cell(50, 6, f"{profit:,}", border=1, align="R", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
